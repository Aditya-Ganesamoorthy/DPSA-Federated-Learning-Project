"""Builds docs/slides.pptx from results/*.json and results/plots/*.png.   python docs/build_slides.py
(needs: pip install python-pptx; dev tooling only, not part of the pipeline)"""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
R, P = ROOT / "results", ROOT / "results" / "plots"
J = lambda n: json.loads((R / n).read_text(encoding="utf-8"))

INK, MUTED, ACCENT = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x2A, 0x78, 0xD6)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(title: str, bullets=None, image=None, image_w=7.4, notes: str = ""):
    s = prs.slides.add_slide(BLANK)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.0))
    p = tb.text_frame.paragraphs[0]
    p.text, p.font.size, p.font.bold, p.font.color.rgb = title, Pt(30), True, INK
    tb.text_frame.word_wrap = True
    if bullets:
        w = 5.0 if image else 12.1
        bx = s.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(w), Inches(5.4))
        tf = bx.text_frame
        tf.word_wrap = True
        for i, b in enumerate(bullets):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text, para.font.size, para.font.color.rgb = "• " + b, Pt(18 if image else 20), INK
            para.space_after = Pt(8)
    if image:
        s.shapes.add_picture(str(P / image), Inches(5.8), Inches(1.5), width=Inches(image_w))
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


cmp_, atk, comm, pc =J("comparison.json"), J("attack.json"), J("communication.json"), J("privacy_curve.json")
m = {x["id"]: x["metrics"] for x in cmp_["methods"]}
auc = lambda k: f"{m[k]['auc']['mean']:.3f}"
e2 = next(e for e in comm["by_local_epochs"] if e["local_epochs"] == 2)

s = prs.slides.add_slide(BLANK)
t = s.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.7), Inches(3))
t.text_frame.word_wrap = True
t.text_frame.paragraphs[0].text = "Decentralized ML: Federated Learning for Privacy Preservation"
t.text_frame.paragraphs[0].font.size, t.text_frame.paragraphs[0].font.bold = Pt(40), True
q = t.text_frame.add_paragraph()
q.text = "Lung-cancer prediction across three simulated hospitals\nFlower · Opacus (DP-SGD) · Secure aggregation · Membership inference"
q.font.size, q.font.color.rgb = Pt(20), MUTED

slide("Problem", ["Centralizing patient records concentrates privacy risk (breach, misuse, regulation)",
                  "Federated learning: train together, keep rows local, share only model updates",
                  "Updates still leak, so we add DP and secure aggregation and measure the cost and the protection"])
slide("Why federated?", ["Each hospital trains locally; the server averages weights: w = Σ (n_k / N) · w_k",
                         "Raw rows never leave a hospital; only HospitalClient code opens a hospital file",
                         "The server evaluates on a shared test file only",
                         "FedAvg verified against a 40-line reference (difference 6e-8)"])
slide("Dataset and privacy preprocessing",
      ["Kaggle lung cancer: 3,000 rows, 15 features, 2 duplicates removed",
       "No identifiers; age → 10-year bands for the report, model uses age/100 (fixed constant)",
       "k-anonymity: pooled 25; per hospital 1–8 (non-IID hospital 3 has k = 1)",
       "Finding: label ≈ unlearnable (CV AUC 0.501 logistic, 0.523 forest): all models near chance"],
      image="clients.png", image_w=7.3)
slide("Architecture and hospitals",
      ["60 / 10 / 20 / 10 split: train / val / test / attack hold-out",
       "Non-IID: 900 / 539 / 359 rows, prevalence 43% / 58% / 58%",
       "MLP 15→32→16→1 (1,057 params); 20 rounds × 2 local epochs; centralized gets a matched 40 epochs",
       "Flower FedAvg + NumPyClient, 5 seeds, one shared test file"])
slide("Methods",
      ["DP-SGD: clip C = 1, noise σ·C, δ = 1e-5, persistent accountant across rounds",
       "Secure aggregation: pairwise masks, integer arithmetic mod 2^32, exact cancellation (simulation)",
       "Membership inference: loss threshold, balanced sets, positive and negative controls",
       "Extras: FedProx, 8/4-bit update quantisation, label-flip poisoning with median defence"])
slide("Live system: Hugging Face + Render + hospital clients",
      ["Primary global model v0 (FedAvg over 3 simulated hospitals) published on Hugging Face (Adhi-gg029/lung-cancer-federated-learning)",
       "Render server = coordinator: login, download model, upload weights, FedAvg into v(k+1), publish; also serves the dashboard",
       "Hospital client: downloads the model, trains on its own CSV (never uploaded), optional local DP, uploads weights + row count",
       "Hospitals and sub-hospitals register themselves; hashed passwords, login throttle, stale-update 409, shape/NaN checks, update-size guard, optional median aggregation",
       "Dashboard: security of every round, raw data never leaves, DP status, attacks and risk score, ordinary-vs-protected demo, EDA",
       "Train & Contribute: upload a CSV batch, train in the browser (optional DP), watch the round animate through to the Hugging Face update",
       "Verified: 99 tests, a real-browser end-to-end round with two hospitals, and a black-box smoke test of the deployment"])
slide("Accuracy: no difference, because the data has no signal",
      [f"Centralized MLP AUC {auc('centralized_mlp')}", f"FedAvg IID {auc('fl_iid')} · non-IID {auc('fl_noniid')}",
       f"FedAvg + SecAgg {auc('fl_secagg')}", f"All inside the 95% chance band 0.5 ± {cmp_['chance_band_auc']:.3f}",
       "Planted-signal check: centralized 0.78, FedAvg 0.79–0.81, FedAvg + DP 0.80"],
      image="comparison_auc.png", image_w=7.2)
slide("Privacy-utility curve",
      ["ε from 69.6 (σ = 0.5) down to 2.9 (σ = 3.0), δ = 1e-5, 20 rounds",
       "No visible utility drop: there is no signal to lose",
       "Smaller hospitals get a larger ε at the same σ"], image="privacy_curve.png", image_w=7.3)
slide("Attack results: leakage falls",
      [f"Centralized {atk['models'][0]['auc']['mean']:.3f} → FedAvg {atk['models'][1]['auc']['mean']:.3f} → "
       f"DP σ=3.0 {atk['models'][-1]['auc']['mean']:.3f}",
       f"Positive control {atk['positive_control']['auc']['mean']:.3f} (attack works)",
       f"Untrained model {atk['untrained_control']['auc']['mean']:.3f} (chance)",
       "One simple attack; low AUC is not proof of privacy"], image="attack.png", image_w=7.3)
slide("Secure aggregation and communication",
      ["Masked integer sum = plain sum (difference 0); a masked update is uncorrelated with the truth",
       "Dropout: masks do not cancel (error ≈ 2,048) → round aborted; real protocols use secret sharing",
       f"FL: {e2['bytes_per_client_per_round'] / 1e3:.1f} KB / hospital / round, {e2['total_mb']:.2f} MB total vs "
       f"{comm['raw_data']['csv_mb']:.3f} MB to upload the raw data once (~6× more)"],
      image="communication.png", image_w=7.0)

ex = R / "extras.json"
if ex.exists():
    X = J("extras.json")
    pl = X["planted_signal_data"]["poisoning"]
    a = lambda d: f"{d['auc']['mean']:.3f}"
    slide("Extras: quantisation, FedProx, poisoning",
          [f"8-bit updates: {X['real_data']['quantization']['8bit']['bytes_up_total_mb'] / X['real_data']['fedavg']['bytes_up_total_mb'] * 100:.0f}% of the upload bytes; "
           f"planted-signal AUC {a(X['planted_signal_data']['quantization']['8bit'])} vs {a(X['planted_signal_data']['fedavg'])}",
           "FedProx: no measurable change on the real data (everything is at chance); see report 6.8",
           f"Label-flip hospital 1 (planted-signal data): FedAvg {a(pl['labelflip_h1_fedavg'])}, median {a(pl['labelflip_h1_median'])} "
           f"(clean {a(pl['clean_fedavg'])})",
           "Poisoning on real labels is unmeasurable: there is no signal to corrupt"])
slide("Limitations",
      ["Near-random survey data; no clinical claim",
       "Simulated (not cryptographic) secure aggregation",
       "One simple attack; in-process Flower; 3 hospitals, ~1.8k rows",
       "Next: real data, real key agreement, shadow-model attacks, robust aggregation"])
slide("Conclusion and demo",
      ["Verified end-to-end FL pipeline: correct DP accounting, exact secure aggregation",
       "Federating and adding DP progressively reduce membership-inference leakage (0.624 → 0.558 → ~0.51)",
       "Accuracy cannot separate methods on this dataset",
       "Demo: dashboard → hospitals → comparison → ε curve → secure-aggregation dropout → one prediction"])

prs.save(ROOT / "docs" / "slides.pptx")
print("saved docs/slides.pptx,", len(prs.slides), "slides")
